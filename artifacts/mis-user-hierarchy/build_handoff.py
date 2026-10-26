from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUTPUT = ROOT / "MIS User Hierarchy Development Handoff.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin, section.bottom_margin = Inches(.66), Inches(.62)
section.left_margin, section.right_margin = Inches(.70), Inches(.70)
section.footer_distance = Inches(.27)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    s = styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
    s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.10
styles['Title'].font.size = Pt(27)
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Subtitle'].font.size = Pt(11)
styles['Subtitle'].font.italic = False
styles['Subtitle'].paragraph_format.space_after = Pt(12)
styles['Heading 1'].font.size = Pt(17)
styles['Heading 1'].paragraph_format.space_before = Pt(0)
styles['Heading 1'].paragraph_format.space_after = Pt(10)
styles['Heading 2'].font.size = Pt(11.5)
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(5)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    styles[name].paragraph_format.keep_with_next = True
for style in styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'MIS User Hierarchy Development Handoff'
doc.core_properties.subject = 'Inebura MIS user hierarchy and data access requirements'
doc.core_properties.author = 'Inebura Product'

def p(text='', style=None, bold_lead=None):
    para = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold=True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para

def h(text, level=1):
    return doc.add_heading(text, level=level)

def bullet(text):
    return p(text, 'List Bullet')

def page(title):
    doc.add_page_break()
    h(title)

def table(headers, rows, widths, font_size=9):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths):
        c.width=Inches(w)
    for i,title in enumerate(headers):
        t.rows[0].cells[i].text=title
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row): cells[i].text=value
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','85'),('bottom','85'),('left','100'),('right','100')]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');margins.append(x)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcpr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'), 'D9E2F3' if ri==0 else ('F4F6F8' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(font_size)
                    run.font.color.rgb=RGBColor(0,0,0)
                    run.bold=ri==0
    p().paragraph_format.space_after=Pt(1)
    return t

def link(para, text, url):
    rid=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'1155CC');rp.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');rp.append(underline)
    r.append(rp);tx=OxmlElement('w:t');tx.text=text;r.append(tx);el.append(r);para._p.append(el)

# Page 1
p('MIS User Hierarchy\nDevelopment Handoff','Title')
p('Inebura   Version 3   23 October 2026','Subtitle')
p('We need to give each MIS User a defined data scope and let an MIS Manager see the combined permitted data of their reporting MIS users. Admin will configure the hierarchy within the existing User Setup journey. Existing reports and downloads must use that hierarchy and the current invoice assignments.')
p('This document includes the revised two-level mappings: selected Agents or Agent cum Managers can be limited by subclient, and selected subclients can be limited by Agents or Agent cum Managers. The requirements apply across Inebura and are broader than a Bluspring-only configuration.')
h('Delivery scope',2)
p('Extend MIS user creation and editing, introduce MIS Manager reporting, apply the selected data scope, and keep report visibility correct through reassignment, suspension, deletion, reactivation and reporting changes. Existing MIS users keep their current access until Admin configures them.')
p('Existing modules, uploads, actions and scheduling already exist. This task changes none of those permissions and does not recreate those features. Data visibility remains separate from permission to perform an action.')
h('Role responsibilities',2)
table(['Role','Requirement'],[
('Admin','Creates and edits MIS users, reporting relationships and data mappings. Admin alone changes this configuration.'),
('MIS User','Reports to exactly one MIS Manager at a time. Views permitted data and downloads reports within the configured scope.'),
('MIS Manager','Views the unique combined data of their reporting MIS users and downloads reports. The role grants no user administration or operational actions.'),
('Agent or Agent cum Manager','Can be selected as the source of personally assigned invoices. Selecting an Agent cum Manager does not include their team\u2019s invoices.')
],[1.35,5.48])
h('Dashboard deferred',2)
p('Dashboard requirements will be updated after obtaining requirements from live clients. The earlier six-card ideas for ODM, Credit, Achievement, Submission, Submission vs Acknowledgement and Issue Tracker remain a reference only. Their metrics, presentation and downloads are not acceptance criteria for this hierarchy delivery.')
p('Cash Posting User, Collection Aging Buckets and the completed Dunning Letter Strategy remain outside this task. INEB-37 is a separate MIS AR-link redirection defect.')
p('Jira handoff will follow review and approval of this document. No Jira ticket has been created for this document revision.')

# Page 2
page('Admin user setup')
h('MIS 01 Retain the existing creation layout',2)
p('Use the live-referenced Create User popup with two columns on desktop and one column on a phone. Keep step one\u2019s identity fields, existing-contact control and Send Login Details toggle. In step two, retain Role Type and Role on the first row, Branch on the next row and Reporting To across the full width. Add MIS mapping controls in this step.')
table(['Control','Required behaviour'],[
('Role Type and Role','Role is disabled until Role Type is selected. Changing Role Type clears Role in the draft. Reporting choices depend on Role. Recognise the existing MIS role and add MIS Manager in the appropriate role group.'),
('Branch','Retain the existing field. Branch does not add a portfolio restriction.'),
('Reporting To','An MIS User selects exactly one eligible MIS Manager. One MIS Manager can have several reporting MIS users.'),
('Access mode','Choose Complete client, Selected users or Selected subclients. Display the relevant primary selector and optional secondary selector.'),
('Complete client','Single-select a valid client. No secondary restriction applies in this mode.'),
('Selected users','Primary: multi-select Agents or Agent cum Managers. Optional secondary: multi-select subclients.'),
('Selected subclients','Primary: multi-select subclients. Optional secondary: multi-select Agents or Agent cum Managers.')
],[1.55,5.28],9)
h('MIS 02 Save one valid scope',2)
p('Complete client is the baseline default. Admin must select a valid client in that mode. New or edited MIS User setup must also have a valid reporting MIS Manager. The other modes require at least one valid primary selection; the secondary restriction is optional.')
p('Primary-only access must be clearly visible as all data matching the primary selections. Refresh dependent secondary choices when Admin changes the primary selections. Remove only truly incompatible or invalid entities. Valid stored entities remain selected even if no invoice currently matches them; such a scope shows no data.')
p('Show the resulting scope for review before save. Do not silently save broader all-data access after a secondary restriction is cleared. Switching mode clears both mapping levels in the Admin draft; the saved access changes only when Admin saves. Edit User uses the same controls and rules.')
h('MIS 03 Preserve existing access during migration',2)
p('Existing MIS users keep current access until Admin explicitly configures them. Do not remove or narrow that access merely because the new fields are initially unconfigured. Once Admin configures and saves a user, apply the saved hierarchy and scope.')

# Page 3
page('Two level data access')
h('MIS 04 Apply the complete configured scope',2)
table(['Mode','Primary selection','Optional secondary restriction'],[
('Complete client','One client. Include every invoice under that client across subclients and assignees.','None. This mode grants the complete selected-client scope.'),
('Selected users','One or more Agents or Agent cum Managers. Include personally assigned invoices across subclients and main clients.','One or more subclients. Include only invoices that match both the selected current assignees and selected subclients.'),
('Selected subclients','One or more subclients. Include their invoices regardless of assignee.','One or more Agents or Agent cum Managers. Include only invoices that match both the selected subclients and selected current assignees.')
],[1.2,2.8,2.83],9)
p('Select one mode at a time. With no secondary restriction, include all invoices matching the primary selections. With both levels selected, use their intersection. Multiple selections within a level are combined; an invoice must match a selected entity at each configured level. A secondary selection can only narrow the primary scope.')
p('Agent cum Manager selections always cover personal assignments. A reporting-team relationship supplies no extra invoices. A valid restricted scope with no matching invoices returns an empty result without dropping its restriction.')
h('Worked examples for both orderings',2)
p('Agent A holds I101 for Client X / North and I102 for Client Y / East. Agent cum Manager B personally holds I103 for Client X / South and I104 for Client X / North. B\u2019s reporting Agent C holds I105 for Client X / North.')
table(['Configured scope','Visible invoices'],[
('Selected users A with no secondary subclients','I101, I102. Agent-only scope crosses clients.'),
('Selected users A with secondary X / North','I101. The secondary subclient excludes I102.'),
('Selected X / North with no secondary users','I101, I104, I105. All assignees in that subclient.'),
('Selected X / North with secondary A','I101. The selected assignee excludes I104 and I105.'),
('Selected Agent cum Manager B with secondary X / North','I104. B\u2019s personal invoice only; C\u2019s I105 is excluded.')
],[3.3,3.53],9)
h('MIS 05 Unique combined manager scope',2)
p('An MIS Manager sees only the union of their active reporting MIS users\u2019 complete configured scopes, including each user\u2019s secondary restrictions. Show and count each invoice once where scopes overlap. The manager role independently grants neither complete-client data nor field-team management.')

# Page 4
page('Assignment and lifecycle rules')
h('MIS 06 Current assignments and reporting changes',2)
p('Reevaluate both configured mapping levels when an invoice\u2019s assignment or subclient changes. An A-only scope loses an invoice moved to unselected B. A scope selecting A and B retains it only if any subclient restriction also matches. A subclient-only scope continues to include that subclient\u2019s invoices regardless of assignee.')
p('Changing an MIS user\u2019s reporting manager updates both managers\u2019 visibility immediately. The new manager gains that user\u2019s permitted data. The former manager retains only data still available through other reporting MIS users. Each MIS user reports to one MIS Manager at a time.')
h('MIS 07 Deleted or suspended assignees',2)
p('Deleting or suspending a selected Agent or Agent cum Manager requires a fallback agent. Transfer A\u2019s affected invoices to fallback B, then replace the affected MIS selection of A with B, whether that selection is primary or secondary. Retain every configured subclient restriction. If B was already selected, retain B once.')
p('Without a subclient restriction, the MIS user sees B\u2019s entire personally assigned portfolio, including B\u2019s existing invoices and those transferred from A. With a subclient restriction, include only B\u2019s personally assigned invoices in the selected subclients. No current matches means an empty result; do not clear a restriction or widen access during fallback.')
p('Example: an MIS user selects A and limits access to X / North. If A falls back to B, the user sees B\u2019s X / North invoices, including matching transfers, but not B\u2019s X / South invoices. The same result applies when X / North is primary and A is the secondary user selection.')
h('MIS 08 Agent reactivation',2)
p('Reactivating A does not restore the earlier MIS selection. Keep fallback B and all retained restrictions until Admin explicitly changes the mappings. Subsequent reassignment follows the complete current scope.')
h('MIS 09 MIS User suspension deletion and reactivation',2)
p('A suspended or deleted MIS user immediately stops contributing scope to their MIS Manager. Another reporting MIS user may still provide access to an overlapping invoice. On reactivation of a suspended MIS user, use their current valid mappings and reporting relationship, including secondary restrictions.')
h('MIS 10 MIS Manager replacement and return',2)
p('When an MIS Manager with reporting MIS users is deleted or suspended, Admin must select a replacement MIS Manager. Move the reporting MIS users to the replacement, retaining their individual mappings. When the former manager returns, those users stay with the replacement until Admin explicitly changes reporting. Do not restore the former hierarchy automatically.')

# Page 5
page('Reports and confirmed clarifications')
h('MIS 11 Apply the current hierarchy at access',2)
p('Both MIS roles can view and download existing reports within their effective scope. Apply the current access mode, every configured mapping level, invoice assignments and reporting relationships when data is viewed and when a download is requested. Existing report filters and date controls retain their behaviour and may only narrow permitted data.')
p('A direct request for a report or file must obey the same scope as the visible controls. MIS Manager report rows and totals use unique combined data; overlapping MIS users must not duplicate or inflate an invoice\u2019s contribution.')
h('MIS 12 Recheck previously generated files',2)
p('An earlier report does not preserve an old access grant. When a stored or previously generated report is retrieved or downloaded later, recheck current scope. Do not return data that is no longer permitted after reassignment, mapping changes, suspension or reporting changes.')
h('MIS 13 Keep operational permissions separate',2)
p('Hierarchy determines whose data is visible. Existing module and action permissions continue to determine what the user can do. This task grants no new upload, customer-maintenance, scheduling, email or other operational permission. It does not recreate existing reporting modules or define new report templates.')
h('Confirmed clarification record',2)
table(['Question','Decision'],[
('Can agent and subclient selections be combined?','Yes. Outside Complete client mode, either can be primary with an optional secondary restriction. Both levels intersect.'),
('Can a secondary selection widen access?','No. It only narrows the selected primary scope.'),
('Does no current invoice match clear a valid restriction?','No. Keep the restriction and show an empty result.'),
('Does fallback B always expose the whole B portfolio?','Only without a subclient restriction. Preserve subclient restrictions in either ordering.'),
('Do existing MIS users lose access before configuration?','No. Keep current access until Admin configures them.'),
('Does MIS User inactivity affect the manager immediately?','Yes. Suspension or deletion removes that user\u2019s contribution.'),
('Do reactivated users and returning managers restore old mappings?','MIS User reactivation uses current valid mappings. A returning MIS Manager\u2019s users stay with the replacement until Admin changes reporting.'),
('Can an older generated report bypass a later scope change?','No. Recheck scope when it is retrieved or downloaded.')
],[2.9,3.93],9)

# Page 6
page('Acceptance scenarios and sources')
p('Verify included and excluded invoices across clients and subclients, personal Agent cum Manager assignments, both mapping orders and overlapping MIS User scopes.')
table(['Scenario','Expected result','Rule'],[
('New or edited setup','Admin saves a valid manager and primary scope. Optional secondary selections are valid and visible for review.','01\u201302'),
('Mode or primary draft change','Dependent options refresh. Mode change clears both draft levels. Saved access remains until save; no silent broadening.','02'),
('Existing user before configuration','Current access continues until Admin configures the user.','03'),
('Complete client','All invoices under the chosen client and none from other clients.','04'),
('Agent primary only','Personal assignments across clients; no team-only invoices.','04'),
('Agent primary and subclient secondary','Only invoices matching selected current assignees and selected subclients.','04'),
('Subclient primary only','Selected subclients\u2019 invoices regardless of assignee.','04'),
('Subclient primary and agent secondary','Only invoices matching both levels; same intersection as the reverse order.','04'),
('Valid restriction with no matches','Empty data; stored valid selections remain unchanged.','02\u201304'),
('MIS Manager view and reporting change','Unique union of complete reporting-user scopes; immediate changes when reporting moves.','05\u201306'),
('Invoice assignment or subclient change','Visibility follows every configured restriction, not only the assignee.','06'),
('Assignee fallback in either order','Replace A with B. Retain subclient restrictions. B\u2019s whole personal portfolio appears only without such a restriction.','07'),
('Agent reactivation','Keep B and retained restrictions until Admin edits them.','08'),
('MIS User inactivity and return','Immediately remove their manager contribution; reactivation uses current valid mappings.','09'),
('MIS Manager replacement and return','Replacement required; returning manager\u2019s former users stay with replacement until Admin changes reporting.','10'),
('Report view download and older file retrieval','Recheck current complete scope. Direct requests and earlier files cannot bypass it; manager totals remain unique.','11\u201312'),
('Existing operational permissions','No module, action, upload, email or scheduling permission changes.','13')
],[1.7,4.43,.70],8.6)
h('Sources and handoff',2)
para=p();link(para,'MIS user hierarchy Space Page','https://chatgpt.com/space/page_18a6b20fa7f081918c5ab739aff81db2')
p('Decisions were recorded in this discussion on 20 and 23 October 2026. The 23 October two-level mapping change supersedes the earlier rule that agent and subclient selections could not intersect. The baseline is Bluspring development requirements for effort analysis, version 2 dated 3 October 2026; later confirmed decisions govern where they refine it.')
para=p('Layout reference inspected on 20 October 2026: ');link(para,'current Inebura user creation popup','https://xeon.inebura.ai/login')
p('Close the document review before creating the Jira ticket. The ticket should cover MIS hierarchy and reporting related to that hierarchy.')

# Footer page field
foot=section.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=foot.add_run('MIS User Hierarchy  |  ');run.font.size=Pt(8);run.font.color.rgb=RGBColor(0,0,0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
fr=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text='1';fr.append(tx);field.append(fr);foot._p.append(field)

doc.save(OUTPUT)
print(OUTPUT)
