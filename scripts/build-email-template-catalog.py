from pathlib import Path
import json,re,html
from html.parser import HTMLParser
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'artifacts/template-master/2026-10-15/templates-source.json').read_text())
class BodyText(HTMLParser):
 def __init__(self):super().__init__(convert_charrefs=True);self.parts=[];self.skip=0;self.links=[]
 def handle_starttag(self,tag,attrs):
  if tag in ['style','script','head']:self.skip+=1;return
  if self.skip:return
  attrs=dict(attrs)
  if tag in ['p','div','h1','h2','h3','h4','li','tr','ul','ol','section']:self.parts.append('\n')
  if tag=='br':self.parts.append('\n')
  if tag=='li':self.parts.append('• ')
  if tag in ['td','th'] and self.parts and not self.parts[-1].endswith('\n'):self.parts.append(' | ')
  if tag=='a':self.links.append(attrs.get('href',''))
 def handle_endtag(self,tag):
  if tag in ['style','script','head']:self.skip=max(0,self.skip-1);return
  if self.skip:return
  if tag in ['p','div','h1','h2','h3','h4','li','tr','section']:self.parts.append('\n')
  if tag=='a' and self.links:
   link=self.links.pop()
   if link and link!='#':self.parts.append(' ('+link+')')
 def handle_data(self,s):
  if not self.skip:self.parts.append(re.sub(r'\s+',' ',s))
 def text(self):
  lines=[]
  for line in ''.join(self.parts).replace('\xa0',' ').splitlines():
   line=re.sub(r'[ \t\r\f\v]+',' ',line).strip(' |')
   if line:lines.append(line)
  return '\n'.join(lines)
def body_text(s):
 if not re.search(r'<(?:!DOCTYPE|html|body|div|p|table|br|span)\b',s,re.I):return s.strip()
 s=re.sub(r'<#--.*?-->', '', s, flags=re.S)
 s=re.sub(r'</?#[^>]*>',lambda m:html.escape(m.group(0)),s)
 parser=BodyText();parser.feed(s);return parser.text()
doc=Document();sec=doc.sections[0];sec.page_height=Inches(11.7);sec.page_width=Inches(8.3);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.05
for name,size in [('Title',25),('Heading 1',19),('Heading 2',13),('Heading 3',10)]:
 st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
# Remove inherited decorative paragraph borders from Word's title styles.
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
# Page numbers support navigation in this long catalog.
f=sec.footer.paragraphs[0];f.alignment=2;r=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
doc.add_paragraph('Inebura Email and WhatsApp Templates','Title')
doc.add_paragraph('Xeon Infraservices Pvt Ltd','Subtitle')
doc.add_paragraph('This catalog contains the message bodies and settings for all 110 templates in Template Master: 78 email templates and 32 WhatsApp templates. Placeholders and conditional template expressions are retained for reference.')
doc.add_paragraph('Captured on 15 October 2026 from Inebura Communication Setup.')
doc.add_paragraph('Failover email','Heading 2')
doc.add_paragraph('Enabled. If none of the configured sender mailboxes can send a message, the message is sent from ar.services@inebura.com. Sign-in emails always use ar.services@inebura.com.')
doc.add_paragraph('Template index','Heading 1')
# Compact index, with the same numbering used in the body.
ordered=[t for mode in ['Email','WhatsApp'] for t in data['templates'] if t['communicationMode']==mode]
for mode in ['Email','WhatsApp']:
 doc.add_paragraph(mode+' templates','Heading 2')
 for i,t in enumerate(ordered,1):
  if t['communicationMode']!=mode:continue
  p=doc.add_paragraph(f'{i:03d}  {t["code"]}');p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
  for r in p.runs:r.font.size=Pt(9)
extracted=[]
for mode in ['Email','WhatsApp']:
 doc.add_page_break();doc.add_paragraph(mode+' templates','Heading 1')
 for i,t in enumerate(ordered,1):
  if t['communicationMode']!=mode:continue
  heading=doc.add_paragraph(f'Template {i:03d}','Heading 2')
  def meta(label,value):
   p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True;p.add_run(label+': ').bold=True;p.add_run(value or '-')
  meta('Code',t['code']);meta('Event type',t['eventType']);meta('Recipient role',t['recipientRole']);meta('Schedule',t['schedule']);meta('Status','Disabled' if t['disabled'] else 'Enabled')
  if t['description']:meta('Description',t['description'])
  if t['subject']:meta('Subject',t['subject'])
  p=doc.add_paragraph('Email body' if mode=='Email' else 'WhatsApp message','Heading 3')
  text=body_text(t['body'])
  for line in text.splitlines():doc.add_paragraph(line)
  notification=body_text(t['notificationBody'])
  if notification and notification!=text:
   doc.add_paragraph('Notification message','Heading 3')
   for line in notification.splitlines():doc.add_paragraph(line)
  extracted.append(dict(number=i,code=t['code'],mode=mode,text=text,notification=notification))
out=ROOT/'outputs/Inebura_Email_and_WhatsApp_Templates.docx';out.parent.mkdir(exist_ok=True);doc.save(out)
(ROOT/'artifacts/template-master/2026-10-15/document-text.json').write_text(json.dumps(extracted,ensure_ascii=False,indent=2))
assert len(extracted)==110 and all(e['text'] for e in extracted)
print(out);print('Entries:',len(extracted),'Words:',sum(len(e['text'].split()) for e in extracted))
